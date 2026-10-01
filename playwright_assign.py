import json
import random
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook, Workbook
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

TODAY = datetime.now().strftime("%Y-%m-%d")
SESSION_DIR = "wa_session"
SCREENSHOT_DIR = Path("screenshots")
SCREENSHOT_DIR.mkdir(exist_ok=True)

NAME_KEYS = ("Name", "Contact Name")
PHONE_KEYS = ("Phone", "Phone Number")
MESSAGE_KEYS = ("Message",)


def pick(record, keys):
    for key in keys:
        if record.get(key):
            return record[key]
    return None


def read_contacts(path="contacts.xlsx"):
    wb = load_workbook(path)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    contacts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        if pick(record, PHONE_KEYS):
            contacts.append(record)
    return contacts


def human_pause(a=2, b=5):
    time.sleep(random.uniform(a, b))


def safe_name(text):
    return re.sub(r'[\\/:*?"<>|\s]+', "_", str(text)).strip("_") or "contact"


def clean_phone(value):
    # keep digits only; handles +91..., spaces, and Excel's 919876543210.0
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"\D", "", text)


def open_chat(page, phone, message):
    """Search for the number in the search box; fall back to a direct link."""
    message_box = page.locator('footer div[contenteditable="true"]').first
    try:
        search_box = page.locator('#side div[contenteditable="true"]').first
        search_box.click()
        search_box.fill(phone)
        page.wait_for_timeout(1500)
        page.locator('#pane-side [role="listitem"], #pane-side [role="row"]').first.click(
            timeout=5000
        )
        message_box.wait_for(state="visible", timeout=8000)
        return "search"
    except PWTimeout:
        # number not in your contacts: open the chat directly
        page.goto(f"https://web.whatsapp.com/send?phone={phone}&text={quote(message)}")
        message_box.wait_for(state="visible", timeout=30000)
        return "link"


def send_message_to_contact(page, contact, report):
    name = str(pick(contact, NAME_KEYS) or "").strip()
    phone = clean_phone(pick(contact, PHONE_KEYS))
    template = pick(contact, MESSAGE_KEYS) or "Hi {name}, this is an automated test message."
    message = str(template).replace("{name}", name)

    entry = {
        "name": name,
        "phone": phone,
        "message": message,
        "status": "failed",
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "last_3_received": [],
        "screenshot": None,
        "error": None,
    }

    try:
        method = open_chat(page, phone, message)
        message_box = page.locator('footer div[contenteditable="true"]').first
        human_pause(1, 2)

        # Outgoing messages have a data-id starting with "true_"
        before = page.locator('div[data-id^="true_"]').count()

        message_box.click()
        if method == "search":
            message_box.fill(message)
        # when opened via link, the text is already pre-filled
        human_pause(1, 2)
        page.keyboard.press("Enter")
        human_pause(3, 4)

        # Screenshot first, so you get one even if the check below fails
        screenshot_path = SCREENSHOT_DIR / f"{safe_name(name or phone)}_{TODAY}.png"
        page.screenshot(path=str(screenshot_path))
        entry["screenshot"] = str(screenshot_path)

        # Confirm a NEW outgoing message appeared
        page.wait_for_function(
            """(before) => document.querySelectorAll('div[data-id^="true_"]').length > before""",
            arg=before,
            timeout=15000,
        )
        entry["status"] = "sent"

        # Smart extraction: last 3 messages received FROM this contact ("false_" = incoming)
        incoming = page.locator('div[data-id^="false_"] span.selectable-text')
        texts = [t.strip() for t in incoming.all_inner_texts() if t.strip()]
        entry["last_3_received"] = texts[-3:]

    except PWTimeout as e:
        entry["error"] = f"Timeout: number invalid or message not confirmed ({e})"
    except Exception as e:
        entry["error"] = str(e)

    report.append(entry)
    human_pause(2, 5)


def save_reports(report):
    json_path = f"whatsapp_report_{TODAY}.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, ensure_ascii=False)

    wb = Workbook()
    ws = wb.active
    ws.append([
        "Name", "Phone", "Status", "Timestamp", "Message",
        "Last 3 Received", "Screenshot", "Error",
    ])
    for entry in report:
        ws.append([
            entry["name"], entry["phone"], entry["status"], entry["timestamp"],
            entry["message"], " | ".join(entry["last_3_received"]),
            entry["screenshot"] or "", entry["error"] or "",
        ])
    xlsx_path = f"whatsapp_report_{TODAY}.xlsx"
    wb.save(xlsx_path)

    print(f"Saved {json_path} and {xlsx_path}")


def main():
    contacts = read_contacts("contacts.xlsx")
    report = []

    try:
        with sync_playwright() as p:
            context = p.chromium.launch_persistent_context(
                SESSION_DIR,
                headless=False,
                channel="chrome",
                args=["--disable-blink-features=AutomationControlled"],
                ignore_default_args=["--enable-automation"],
            )
            page = context.pages[0] if context.pages else context.new_page()
            page.goto("https://web.whatsapp.com")

            print("Scan the QR code in the Chrome window (up to 5 minutes)...")
            page.wait_for_selector("#pane-side", timeout=300000)
            print("Logged in. Starting message run.")

            for contact in contacts:
                send_message_to_contact(page, contact, report)

            context.close()
    finally:
        save_reports(report)


if __name__ == "__main__":
    main()